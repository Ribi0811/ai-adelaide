from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

OUT=Path(__file__).parent
D=Document(); sec=D.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.67)
sec.left_margin=sec.right_margin=Inches(.78)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 s=D.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_after=Pt(7)
D.styles['Normal'].font.size=Pt(11)
D.styles['Normal'].paragraph_format.line_spacing=1.10
D.styles['Title'].font.size=Pt(29);D.styles['Title'].font.bold=True
D.styles['Title'].paragraph_format.space_after=Pt(12)
D.styles['Subtitle'].font.size=Pt(11)
D.styles['Heading 1'].font.size=Pt(19);D.styles['Heading 1'].font.bold=True
D.styles['Heading 2'].font.size=Pt(13);D.styles['Heading 2'].font.bold=True
D.styles['Heading 2'].paragraph_format.space_before=Pt(12)
for node in list(D.styles.element.iter(qn('w:pBdr'))):
 node.getparent().remove(node)
for node in list(D.styles.element.iter(qn('w:contextualSpacing'))):
 node.getparent().remove(node)
D.core_properties.author='AI Adelaide';D.core_properties.title='My Dezign Exteriors website proposal'
D.core_properties.subject='Website redesign options and next steps for mdex.com.au'
D.core_properties.keywords=''

def p(t='',style=None):return D.add_paragraph(t,style)
def h(t):return D.add_heading(t,2)
def lead(label,t):
 a=p();a.add_run(label+' ').bold=True;a.add_run(t);return a
def page():D.add_page_break()
def table(headers,rows,widths):
 t=D.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for c,s in zip(t.rows[0].cells,headers):c.text=s
 for row in rows:
  for c,s in zip(t.add_row().cells,row):c.text=s
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'5');b.set(qn('w:color'),'D9D9D9');borders.append(b)
 t._tbl.tblPr.append(borders)
 for i,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for c,w in zip(row.cells,widths):
   c.width=Inches(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    v=OxmlElement('w:'+side);v.set(qn('w:w'),'110');v.set(qn('w:type'),'dxa');m.append(v)
   tcpr.append(m);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'253743' if i==0 else ('F1F4F5' if i%2 else 'FFFFFF'));tcpr.append(shade)
   for a in c.paragraphs:
    a.paragraph_format.space_after=Pt(2);a.paragraph_format.line_spacing=1.04
    for r in a.runs:
     r.font.size=Pt(10);r.font.bold=(i==0);r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
 p().paragraph_format.space_after=Pt(0)
 return t

p('AI ADELAIDE   /   PREPARED FOR MY DEZIGN EXTERIORS','Subtitle')
D.add_heading('My Dezign Exteriors\nwebsite proposal',0)
p('Design direction and options for mdex.com.au  |  6 September 2026','Subtitle')
p('We recommend a website that puts your completed outdoor spaces at the centre, helps homeowners understand their options and makes it easier to enquire. We would preserve your existing web address and useful page links, and make recent work easier to keep up to date.')
h('Build on what you already have')
p('Your site already has service galleries, customer reviews, supplier information and forms that accept plans or photos. These are useful foundations. The opportunity is to give your work more room and make the site easier to explore, especially on a phone.')
p('Our review found a large form occupying much of the opening screen and a mobile image crop that shows mostly roofing. The footer FAQ link also needs correcting. Your Facebook project feed is already connected, although several images did not load during our visit; we would check its setup before deciding whether to retain it.')
h('The proposed look and experience')
lead('Lead with a finished space.', 'Use your strongest original project photography, generous type and a clean charcoal and warm-white palette, with your existing blue as an accent. We would confirm image ownership and project details with you.')
lead('Help customers find the right work.', 'Feature a few projects with the customer’s brief, the solution and key materials. Let visitors explore verandahs, carports, outdoor rooms, decks and screens. Use a before-and-after slider where matching photos are available.')
lead('Make the next step obvious.', 'Offer clear View our work and Discuss your project buttons, easy phone access and a shorter enquiry form. Keep photo and plan uploads optional.')
h('A first concept to review')
p('We suggest starting with a homepage concept and one service or project page. This would show the visual direction before the full site is built. The final scope, fee and schedule would be agreed before work begins.')

page()
D.add_heading('How you want to manage the site',1)
p('The most useful choice is how involved you want to be in updating the website. Both WordPress and Next.js can deliver a polished, mobile-friendly result. Neither automatically produces better Google rankings.')
table(['','Refreshed WordPress','Next.js with a simple editor'],[
 ('Best suited to','You want to keep a familiar dashboard and edit pages yourself.','You want us to manage the website, with straightforward project updates.'),
 ('What changes','We redesign the site and review the existing theme, plugins and feeds.','We rebuild the public site and provide an agreed way to edit project content.'),
 ('Ongoing care','Hosting, backups, plugin and theme updates, and form checks.','Hosting, application updates, content backups, integrations and form checks.'),
 ('Main tradeoff','Less migration work, with continued plugin maintenance.','More migration work; editing features must be included in the scope.')
],[1.00,2.855,2.855])
p('Our recommendation: keep WordPress if familiar self-editing matters most. If you prefer us to look after the site, we would explore Next.js with a small project editor. We would confirm this after understanding how you currently make changes.')
h('Keep sharing your work on social media')
p('You could continue posting to Facebook or Instagram. A connected feed could display those updates in a Latest updates section, alongside a separate gallery of selected projects. Your main portfolio would remain available if a social connection stops working.')
p('Your current Recent Projects page already uses a Facebook feed. We would check the account connection, subscription and image loading first. Instagram support would depend on your account type and the selected integration.')
h('An optional workflow for later')
p('A new social post could create a website project draft. You confirm details such as the suburb and materials, then approve a full project story and a tailored Google Business Profile update. This would be a separately scoped option after the website launch.')
p('We would confirm account access, supported post types and running costs before offering this automation. Google Business Profile posts appear on the business profile; they do not guarantee website indexing or higher rankings.')

page()
D.add_heading('A careful switch and clear ongoing costs',1)
h('Protect the existing website value')
p('The redesign would keep mdex.com.au and retain useful paths such as /verandahs/, /carports/ and /recent-projects/. Before any switch, we would inventory pages, images, downloads and existing redirects, and review available Google Search Console data.')
p('Where an address needs to change, we would map it to the closest relevant page. We would also check mobile layouts, forms and uploads, search settings and enquiry delivery. The launch plan would protect your email setup and retain a recovery copy of the old site. Search positions can fluctuate during a migration.')
h('Agree the commercial details first')
p('We would provide a written scope showing the one-off redesign and migration fee, plus a separate monthly or annual care price. Any neighbour discount would be stated clearly in that quote. There is no commitment to a build or subscription in this outline.')
p('Care could cover hosting, backups, software updates, form checks and an agreed allowance for small edits. New features, regular SEO work and custom social automation would be itemised separately. Domain, email and third-party subscription charges would also be clear, together with ownership and handover terms.')
h('A few things we would like to understand')
for label,txt in [
 ('Your ideal enquiries','Which jobs would you like more of, and which suburbs or areas do you serve?'),
 ('Your editing preference','Would you prefer to update the site yourself, have us handle changes, or do a bit of both? Would you like to keep WordPress?'),
 ('Your recent work','Do you have original photos or videos of three to five projects we could feature, with permission to use them?'),
 ('Your social updates','Where do you normally post, and would you prefer an automatic feed or approval of selected website stories?'),
 ('Your current arrangement','What do you currently pay for hosting, email, domain and maintenance? A renewal invoice would help us compare the inclusions.'),
 ('Your budget and timing','What budget range and launch timing would feel comfortable? Who should receive website enquiries?')]:
 lead(label+'.',txt)
p('You can reply in a few lines or talk these through with us. We can then recommend the right approach and prepare a clear quote. Please do not send passwords; any access would be arranged securely when needed.')
p('AI Adelaide  |  aiadelaide.com.au','Subtitle')
D.save(OUT/'MDEX-Website-Proposal.docx')
print(OUT/'MDEX-Website-Proposal.docx')
